"""
Escreve os SlideSpec num arquivo .pptx a partir do TEMPLATE.pptx.
"""
from __future__ import annotations

from typing import Optional

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide, SlideLayout

from content_model.sunday_mass import SundayMass
from pptx_builder.deck_plan import SlideSpec, build_full_deck

TITLE_LAYOUT_NAME = "PASCOM_Título"
CONTENT_LAYOUT_NAME = "PASCOM_Título e Conteúdo"

TITLE_PLACEHOLDER_IDX = 0
CONTENT_PLACEHOLDER_IDX = 1

# Precisa bater com _ROLE_FONT_VARIANT em deck_plan.py (que escolhe a fonte
# usada pra *medir* o texto) — aqui é a fonte aplicada de fato no slide.
_ROLE_FONT = {
    "leader": ("Arial", False),
    "assembly": ("Arial", True),
}


def _find_layout(prs: PresentationType, name: str) -> SlideLayout:
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    raise ValueError(f"Layout {name!r} não encontrado no template.")


def _delete_slide(prs: PresentationType, index: int) -> None:
    """python-pptx não tem uma API nativa pra remover slide; a técnica
    padrão é remover a entrada correspondente de sldIdLst diretamente."""
    xml_slides = prs.slides._sldIdLst
    slide_id_elements = list(xml_slides)
    xml_slides.remove(slide_id_elements[index])


def _set_placeholder_text(
    slide: Slide, idx: int, paragraphs: list[str], *, role: Optional[str] = None
) -> None:
    placeholder = slide.placeholders[idx]
    text_frame = placeholder.text_frame
    text_frame.text = paragraphs[0]
    for extra in paragraphs[1:]:
        p = text_frame.add_paragraph()
        p.text = extra

    font_override = _ROLE_FONT.get(role)
    if font_override is not None:
        font_name, bold = font_override
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = font_name
                run.font.bold = bold


def render_slide_specs(prs: PresentationType, specs: list[SlideSpec]) -> None:
    title_layout = _find_layout(prs, TITLE_LAYOUT_NAME)
    content_layout = _find_layout(prs, CONTENT_LAYOUT_NAME)

    for spec in specs:
        if spec.layout == "title":
            slide = prs.slides.add_slide(title_layout)
            _set_placeholder_text(slide, TITLE_PLACEHOLDER_IDX, [spec.title or ""])
        elif spec.layout == "content":
            slide = prs.slides.add_slide(content_layout)
            _set_placeholder_text(slide, TITLE_PLACEHOLDER_IDX, [spec.title or ""])
            _set_placeholder_text(
                slide, CONTENT_PLACEHOLDER_IDX, spec.body or [""], role=spec.role
            )
        else:
            raise ValueError(f"Layout de slide desconhecido: {spec.layout!r}")


def build_pptx(mass: SundayMass, template_path: str, output_path: str) -> None:
    prs = Presentation(template_path)

    # As 2 slides que já vêm no template são só exemplo de autoria de cada
    # layout — removemos antes de gerar o deck real.
    seed_slide_count = len(prs.slides)

    content_layout = _find_layout(prs, CONTENT_LAYOUT_NAME)
    content_placeholder = content_layout.placeholders[CONTENT_PLACEHOLDER_IDX]
    box_width_emu = content_placeholder.width
    box_height_emu = content_placeholder.height

    specs = build_full_deck(
        mass, box_width_emu=box_width_emu, box_height_emu=box_height_emu
    )
    render_slide_specs(prs, specs)

    for i in range(seed_slide_count - 1, -1, -1):
        _delete_slide(prs, i)

    prs.save(output_path)
